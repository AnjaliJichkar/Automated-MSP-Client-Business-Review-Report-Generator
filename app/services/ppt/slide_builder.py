"""
Slide Builder — helpers to create slides, headers, footers, and place charts.
"""
import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

import app.services.ppt.layout_manager as lm



def _blank_slide(prs):
    """Add a blank slide (layout 6) with white background."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = lm.CLR_SLIDE_BG
    return slide



def add_header(slide, section_name: str, slide_title: str = '', font_size: int = 25, header_height_in: float = None):
    """
    Dark green header bar at the top of the slide.
    Left side: "SectionName  |  Slide Title"
    """
    h = header_height_in if header_height_in is not None else lm.HEADER_H_IN
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(lm.SLIDE_W_IN), Inches(h)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = lm.HEADER_BG
    bar.line.fill.background()

    tf = bar.text_frame
    tf.margin_left    = Inches(0.2)
    tf.margin_top     = Inches(0)
    tf.margin_bottom  = Inches(0)
    tf.margin_right   = Inches(0.1)
    tf.word_wrap      = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT

    run_section = p.add_run()
    run_section.text = section_name
    run_section.font.name  = lm.FONT_NAME
    run_section.font.size  = Pt(font_size)
    run_section.font.bold  = True
    run_section.font.color.rgb = lm.HEADER_TEXT_CLR

    if slide_title:
        run_sep = p.add_run()
        run_sep.text = "   |   "
        run_sep.font.size = Pt(16)
        run_sep.font.color.rgb = RGBColor(0xA7, 0xD7, 0xBE)

        run_title = p.add_run()
        run_title.text = slide_title
        run_title.font.name  = lm.FONT_NAME
        run_title.font.size  = Pt(16)
        run_title.font.bold  = False
        run_title.font.color.rgb = lm.HEADER_TEXT_CLR



def add_footer(slide):
    """Add footer_banner.png at the bottom of the slide."""
    if os.path.exists(lm.IMG_FOOTER_BANNER):
        slide.shapes.add_picture(
            lm.IMG_FOOTER_BANNER,
            lm.FOOTER_LEFT, lm.FOOTER_TOP,
            lm.FOOTER_WIDTH, lm.FOOTER_HEIGHT
        )



def add_cover_slide(prs, org_name: str, company: str, period: str, generated_on: str = ''):
    """
    Cover slide matching the Client Business Review template layout.
    """
    slide = _blank_slide(prs)

    if os.path.exists(lm.IMG_TEAMLOGIC_LOGO):
        slide.shapes.add_picture(
            lm.IMG_TEAMLOGIC_LOGO,
            lm.COVER_LOGO_LEFT, lm.COVER_LOGO_TOP,
            lm.COVER_LOGO_W,    lm.COVER_LOGO_H
        )

    display_name = org_name or company or ''
    if display_name:
        org_box = slide.shapes.add_textbox(
            Inches(3.7578), Inches(1.3033), Inches(2.4844), Inches(0.3702)
        )
        tf = org_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = display_name
        run.font.name   = lm.FONT_NAME
        run.font.size   = Pt(16)
        run.font.bold   = True
        run.font.italic = True
        run.font.color.rgb = lm.CLR_COVER_TITLE

    title_box = slide.shapes.add_textbox(
        Inches(3.096), Inches(1.72), Inches(3.808), Inches(0.5385)
    )
    tf2 = title_box.text_frame
    tf2.word_wrap = False
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    report_title = (
        "Quarterly Report"
        if period.strip().lower().startswith("last 3")
        else "Monthly Report"
    )
    run2.text = report_title
    run2.font.name  = lm.FONT_NAME
    run2.font.size  = Pt(26)
    run2.font.bold  = True
    run2.font.color.rgb = lm.CLR_COVER_TITLE

    divider = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(2.5), Inches(2.55), Inches(5.0), Inches(0.03)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = lm.HEADER_BG
    divider.line.fill.background()

    info_labels = ["Presented to:", "Presented by:"]
    info_values = [display_name, company or "TeamLogic IT"]
    if period:
        info_labels.append("Report Period:")
        info_values.append(period)
    if generated_on:
        info_labels.append("Generated On:")
        info_values.append(generated_on)

    # compute label and value boxes so values align exactly on the same x-axis
    lbl_left = 3.7578
    lbl_w = 1.103
    gap = 0.10
    val_left = lbl_left + lbl_w + gap

    lbl_box = slide.shapes.add_textbox(
        Inches(lbl_left), Inches(2.6467), Inches(lbl_w), Inches(1.1518)
    )
    lbl_tf = lbl_box.text_frame
    lbl_tf.word_wrap = False
    lbl_tf.margin_left = Inches(0)
    lbl_tf.margin_right = Inches(0)
    lbl_tf.margin_top = Inches(0)
    lbl_tf.margin_bottom = Inches(0)
    lbl_tf.vertical_anchor = MSO_ANCHOR.TOP

    val_box = slide.shapes.add_textbox(
        Inches(val_left), Inches(2.6467), Inches(1.7241), Inches(1.1518)
    )
    val_tf = val_box.text_frame
    # keep wrapping off so value lines line up with labels vertically
    val_tf.word_wrap = False
    val_tf.margin_left = Inches(0)
    val_tf.margin_right = Inches(0)
    val_tf.margin_top = Inches(0)
    val_tf.margin_bottom = Inches(0)
    val_tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, (lbl, val) in enumerate(zip(info_labels, info_values)):
        lp = lbl_tf.paragraphs[0] if i == 0 else lbl_tf.add_paragraph()
        lp.alignment = PP_ALIGN.LEFT
        lp.space_after = Pt(5)
        lr = lp.add_run()
        lr.text = lbl
        lr.font.name = lm.FONT_NAME
        lr.font.size = Pt(11)
        lr.font.color.rgb = lm.CLR_COVER_SUBTITLE

        vp = val_tf.paragraphs[0] if i == 0 else val_tf.add_paragraph()
        vp.alignment = PP_ALIGN.LEFT
        vp.space_after = Pt(5)
        vr = vp.add_run()
        vr.text = val
        vr.font.name = lm.FONT_NAME
        vr.font.size = Pt(11)
        vr.font.color.rgb = lm.CLR_COVER_SUBTITLE

    copy_box = slide.shapes.add_textbox(
        Inches(3.4107), Inches(5.05), Inches(3.1787), Inches(0.2188)
    )
    cp = copy_box.text_frame.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run()
    cr.text = "\u00a9 2019-2026 TeamLogic, Inc. All rights reserved.  |  Business Confidential"
    cr.font.name  = lm.FONT_NAME
    cr.font.size  = Pt(7)
    cr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    if os.path.exists(lm.IMG_FOOTER_BANNER):
        slide.shapes.add_picture(
            lm.IMG_FOOTER_BANNER,
            lm.COVER_FOOTER_LEFT, lm.COVER_FOOTER_TOP,
            lm.COVER_FOOTER_W,    lm.COVER_FOOTER_H
        )



def add_section_slide(prs, section_name: str, subtitle: str = ''):
    """
    Full-slide dark green section divider with integration name centered.
    """
    slide = _blank_slide(prs)

    # Dark background
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = lm.CLR_SECTION_BG

    # Large section name
    tb = slide.shapes.add_textbox(



        Inches(1.0), Inches(1.6), Inches(8.0), Inches(1.1)
    )
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = section_name
    run.font.name  = lm.FONT_NAME
    run.font.size  = Pt(32)
    run.font.bold  = True
    run.font.color.rgb = lm.CLR_WHITE

    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.name  = lm.FONT_NAME
        r2.font.size  = Pt(13)
        r2.font.bold  = False
        r2.font.color.rgb = RGBColor(0xA7, 0xD7, 0xBE)

    # Thin accent line below text
    line_box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(3.0), Inches(2.6), Inches(4.0), Inches(0.04)
    )
    line_box.fill.solid()
    line_box.fill.fore_color.rgb = RGBColor(0xA7, 0xD7, 0xBE)
    line_box.line.fill.background()

    # Footer
    add_footer(slide)
    return slide



def new_content_slide(prs, section_name: str, slide_title: str = ''):
    """
    Create a blank white content slide with header + footer.
    Returns the slide object ready for chart placement.
    """
    slide = _blank_slide(prs)
    add_header(slide, section_name, slide_title)
    add_footer(slide)
    return slide



def left_slot():
    """Returns (left_in, top_in, w_in, h_in) for the LEFT chart in a 2-chart layout."""
    return (lm.CHART2_L1_IN, lm.CHART2_TOP_IN, lm.CHART2_W_IN, lm.CHART2_H_IN)


def right_slot():
    """Returns (left_in, top_in, w_in, h_in) for the RIGHT chart in a 2-chart layout."""
    return (lm.CHART2_L2_IN, lm.CHART2_TOP_IN, lm.CHART2_W_IN, lm.CHART2_H_IN)


def full_slot():
    """Returns (left_in, top_in, w_in, h_in) for a single full-width chart."""
    return (lm.CHART1_LEFT_IN, lm.CHART1_TOP_IN, lm.CHART1_W_IN, lm.CHART1_H_IN)


def build_dynamic_slide(prs, section_name: str, slide_title: str, chart_specs: list):
    """
    Build a content slide for all chart specs.
    chart_specs: list of (has_data_fn, render_fn) tuples.
      has_data_fn()          → bool — used to decide whether this chart gets
                                       a slot at all. A chart whose data was
                                       filtered out (deselected) or is simply
                                       missing is skipped entirely — no empty
                                       donut, no empty slot.
      render_fn(slide, l, t, w, h) → None — draws the chart into the given slot

    Only specs whose has_data_fn() is True get a slot; slots are then divided
    evenly across just those. If NO spec has data, the slide itself is never
    created — an integration section with nothing active shows no slide.

    render_fn() is still wrapped in try/except as a safety net: if a chart
    that DOES have data still fails to render for some other reason, we fall
    back to the plain placeholder so the slide never breaks.
    """
    from app.services.ppt.chart_builder import add_no_data

    if not chart_specs:
        return None

    active_specs = [spec for spec in chart_specs if spec[0]()]
    if not active_specs:
        return None

    slide = new_content_slide(prs, section_name, slide_title)
    slots = lm.calc_slots(len(active_specs))
    for spec, slot in zip(active_specs, slots):
        has_fn, render_fn = spec[0], spec[1]
        chart_title = spec[2] if len(spec) > 2 else "No Data Available"
        try:
            render_fn(slide, slot[0], slot[1], slot[2], slot[3])
        except Exception:
            add_no_data(slide, chart_title, slot[0], slot[1], slot[2], slot[3])
    return slide